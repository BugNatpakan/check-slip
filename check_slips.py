#!/usr/bin/env python3
"""LG26 payment slip pre-screener.

Offline by default. If [slip2go] enabled = true in config.toml, new slips are also
checked online with Slip2Go (only the QR text is sent, each slip once).

Run with no arguments to process the latest zip in inbox/:
    python check_slips.py
Options:
    --zip PATH        use a specific zip instead of the latest one
    --redecode        ignore the decode cache and decode every file again
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import os
import random
import re
import shutil
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

try:
    import tomllib
except ModuleNotFoundError:  # Python < 3.11
    import tomli as tomllib

try:
    import cv2
    import imagehash
    import numpy as np
    import openpyxl
    import zxingcpp
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from PIL import Image, ImageOps
    from tqdm import tqdm

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
except ImportError as e:
    sys.exit(f"Missing package: {e.name}. Activate the venv and run: pip install -r requirements.txt")

BASE = Path(__file__).resolve().parent
INBOX = BASE / "inbox"
OUTPUT = BASE / "output"
HISTORY = OUTPUT / "history"
CACHE = BASE / ".cache"
EXTRACTED = CACHE / "extracted"
STATE_FILE = CACHE / "state.json"
DECODE_CACHE_FILE = CACHE / "decode_cache.json"
API_CACHE_FILE = CACHE / "slip2go.json"
SECRET_FILE = BASE / "slip2go_secret.txt"
XLSX = OUTPUT / "slip_check.xlsx"
CSV_OUT = OUTPUT / "slip_check.csv"

DECODER_VERSION = 2
ZIP_DATE_RE = re.compile(r"payment-slips-(\d{4}-\d{2}-\d{2})\.zip$", re.I)
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
PDF_DPI = 200
PDF_MAX_PAGES = 5
THAI_TZ = dt.timezone(dt.timedelta(hours=7))

BANKS = {
    "002": "BBL", "004": "KBANK", "006": "KTB", "011": "TTB", "014": "SCB",
    "022": "CIMB", "024": "UOB", "025": "BAY", "030": "GSB", "033": "GHB",
    "034": "BAAC", "067": "TISCO", "069": "KKP", "073": "LHBANK",
}

# Flags in display order. NOT_PAID is a status, not a problem, and is kept apart.
MONEY_FLAGS = ["WRONG_RECEIVER", "SLIP_FRAUD", "SLIP_NOT_FOUND", "UNDERPAID", "OVERPAID"]
FLAG_ORDER = ["NO_QR", "NOT_SLIP_QR", "BAD_CRC", *MONEY_FLAGS, "WRONG_DATE", "DUPLICATE_REF",
              "DUPLICATE_IMAGE", "FILE_MISSING", "SLIP_CHANGED", "API_PENDING"]
NOT_PAID = "NOT_PAID"

VERIFIED, PROBLEM, PENDING = "✅ Verified", "❌ Problem", "⏳ Pending"
MANUAL_COLS = ["manual_status", "checked_by", "note"]
COLUMNS = ["lgNumber", "firstName", "lastName", "fileStatus", "file", "slip_count", "amount_paid",
           "amount_detail", "payment_status", "receiver_check", "transfer_time", "sender_name",
           "bank_name", "transaction_ref", "crc_ok", "api_result", "decode_method", "flags",
           "in_sample", "first_seen", *MANUAL_COLS]
SLIP_COLUMNS = ["lgNumber", "firstName", "page", "bank_name", "transaction_ref", "crc_ok", "amount",
                "transfer_time", "sender_name", "receiver_name", "receiver_account", "receiver_check",
                "api_code", "api_message", "checked_at"]

FILL = {
    "red": PatternFill("solid", fgColor="FFC7CE"),
    "yellow": PatternFill("solid", fgColor="FFEB9C"),
    "green": PatternFill("solid", fgColor="A9D08E"),
    "lightgreen": PatternFill("solid", fgColor="E2F0D9"),
    "grey": PatternFill("solid", fgColor="D9D9D9"),
}
PASSBOOK_LABEL = "passbook total (fill in)"

# Slip2Go result codes -> short label shown in api_result
API_OK_CODES = {"200000", "200200", "200501"}  # 200501 = Slip2Go's own duplicate check; ours is local
API_LABELS = {
    "200401": "WRONG RECEIVER", "200402": "AMOUNT MISMATCH", "200403": "WRONG DATE",
    "200404": "NOT FOUND", "200500": "FRAUD", "400001": "INVALID QR", "400400": "INVALID REQUEST",
}
API_RETRY_CODES = {"200502", "429000", "500500", "500503", "NETWORK", "BADRESPONSE"}
API_AUTH_MESSAGES = {
    "401001": "secret key is wrong", "401002": "shop not found", "401003": "account suspended",
    "401004": "package expired", "401005": "tokens used up", "401006": "credit too low",
    "401007": "this IP address is not allowed (check the IP whitelist)",
}


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s or "")


def status_kind(s: str | None) -> str:
    """Map a manual_status cell to verified/problem/pending, tolerating hand typing."""
    s = (s or "").lower()
    if "verif" in s or "✅" in s:
        return "verified"
    if "problem" in s or "❌" in s:
        return "problem"
    return "pending"


def baht(x: float) -> str:
    return f"{x:,.0f}" if float(x).is_integer() else f"{x:,.2f}"


# ---------------------------------------------------------------- slip payload

def crc16_ccitt_false(data: bytes) -> int:
    crc = 0xFFFF
    for b in data:
        crc ^= b << 8
        for _ in range(8):
            crc = ((crc << 1) ^ 0x1021) if crc & 0x8000 else (crc << 1)
            crc &= 0xFFFF
    return crc


def parse_tlv(s: str):
    """Parse a 2-digit-tag / 2-digit-length TLV string. Returns [(tag, value, start)] or None."""
    out, i = [], 0
    while i < len(s):
        if i + 4 > len(s) or not s[i:i + 4].isdigit():
            return None
        tag, length = s[i:i + 2], int(s[i + 2:i + 4])
        if i + 4 + length > len(s):
            return None
        out.append((tag, s[i + 4:i + 4 + length], i))
        i += 4 + length
    return out or None


def parse_slip(text: str):
    """Return slip fields if `text` is a Thai bank slip verification QR, else None."""
    items = parse_tlv(text.strip())
    if not items:
        return None
    top = {t: (v, pos) for t, v, pos in items}
    if "00" not in top:
        return None
    sub = parse_tlv(top["00"][0])
    if not sub:
        return None
    sd = {t: v for t, v, _ in sub}
    if sd.get("00") != "000001" or "01" not in sd or not sd.get("02"):
        return None
    if "51" in top and top["51"][0] != "TH":
        return None
    payload = text.strip()
    crc_ok = False
    if "91" in top and len(top["91"][0]) == 4:
        pos = top["91"][1]
        calc = crc16_ccitt_false(payload[:pos + 4].encode("ascii", "replace"))
        crc_ok = f"{calc:04X}" == top["91"][0].upper()
    code = sd["01"]
    return {"bank_code": code, "bank_name": BANKS.get(code, f"UNKNOWN({code})"),
            "ref": sd["02"], "crc_ok": crc_ok}


def describe_other_qr(text: str) -> str:
    t = text.strip()
    if t.startswith("000201"):
        return "PromptPay/EMV payment QR"
    if re.match(r"^[a-z]+://", t, re.I):
        return "URL"
    return "unknown QR"


# ---------------------------------------------------------------- QR decoding

def _zxing(img) -> list[str]:
    try:
        return [r.text for r in zxingcpp.read_barcodes(img, formats=zxingcpp.BarcodeFormat.QRCode) if r.text]
    except Exception:
        return []


_CV_DETECTOR = None


def _opencv(img) -> list[str]:
    global _CV_DETECTOR
    try:
        arr = np.asarray(img)
        if arr.ndim == 3:
            arr = cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)
        if _CV_DETECTOR is None:
            _CV_DETECTOR = cv2.QRCodeDetector()
        ok, texts, _, _ = _CV_DETECTOR.detectAndDecodeMulti(arr)
        return [t for t in texts if t] if ok else []
    except Exception:
        return []


def _scaled(img, factor: float, max_side: int = 5000):
    w, h = img.size
    factor = min(factor, max_side / max(w, h))
    if factor <= 1.01:
        return img
    return img.resize((int(w * factor), int(h * factor)), Image.LANCZOS)


def _adaptive(img):
    arr = np.asarray(img.convert("L"))
    return Image.fromarray(cv2.adaptiveThreshold(
        arr, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 51, 10))


def _crop_boxes(w, h):
    ox, oy = int(w * 0.1), int(h * 0.1)
    return {
        "top": (0, 0, w, h // 2 + oy), "bottom": (0, h // 2 - oy, w, h),
        "left": (0, 0, w // 2 + ox, h), "right": (w // 2 - ox, 0, w, h),
        "q-tl": (0, 0, w // 2 + ox, h // 2 + oy), "q-tr": (w // 2 - ox, 0, w, h // 2 + oy),
        "q-bl": (0, h // 2 - oy, w // 2 + ox, h), "q-br": (w // 2 - ox, h // 2 - oy, w, h),
    }


def _variants(img):
    """Preprocessed versions of an image, cheapest first."""
    gray = img.convert("L")
    up = _scaled(gray, 2)
    yield "gray", gray
    yield "upscale2x", up
    yield "adaptive", _adaptive(up)
    for name, box in _crop_boxes(*gray.size).items():
        crop = _scaled(gray.crop(box), 2)
        yield f"crop-{name}", crop
        yield f"crop-{name}-adaptive", _adaptive(crop)
    for angle in (-3, 3, -7, 7, -12, 12):
        yield f"rot{angle}", gray.rotate(angle, expand=True, fillcolor=255, resample=Image.BICUBIC)


def decode_pages(pages: list) -> dict:
    """Find every slip QR on every page.

    Returns dict(slips=[{payload, page, method, count}], others=[non-slip QR texts]).
    `count` is how many times the same slip QR appears in the file (2+ = pasted twice).
    """
    slips: dict[str, dict] = {}
    others: list[str] = []

    def take(texts, page_no, method, page_counts):
        hit = False
        per_call: dict[str, int] = {}
        for t in texts:
            t = t.strip()
            if parse_slip(t):
                hit = True
                per_call[t] = per_call.get(t, 0) + 1
                if t not in slips:
                    slips[t] = {"payload": t, "page": page_no, "method": method, "count": 0}
            elif t not in others:
                others.append(t)
        for t, n in per_call.items():
            page_counts[t] = max(page_counts.get(t, 0), n)
        return hit

    multi = len(pages) > 1
    for pi, page in enumerate(pages, 1):
        prefix = f"p{pi} " if multi else ""
        page_counts: dict[str, int] = {}
        found = take(_zxing(page), pi, prefix + "zxing", page_counts)
        found |= take(_opencv(page), pi, prefix + "opencv", page_counts)
        # crops catch a second slip pasted into the same image
        for name, box in _crop_boxes(*page.size).items():
            found |= take(_zxing(page.crop(box)), pi, f"{prefix}zxing+crop-{name}", page_counts)
        if not found:
            for name, var in _variants(page):
                if take(_zxing(var), pi, f"{prefix}zxing+{name}", page_counts) or \
                        take(_opencv(var), pi, f"{prefix}opencv+{name}", page_counts):
                    break
        for t, n in page_counts.items():
            slips[t]["count"] += n
    return {"slips": list(slips.values()), "others": others}


def load_pages(path: Path) -> list:
    if path.suffix.lower() == ".pdf":
        pages = []
        with fitz.open(path) as doc:
            for page in list(doc)[:PDF_MAX_PAGES]:
                pix = page.get_pixmap(dpi=PDF_DPI)
                pages.append(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))
        return pages
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[-1])
            return [bg]
        return [im.convert("RGB")]


def analyse_file(path: Path) -> dict:
    """Decode + perceptual hash for one file. Never raises."""
    entry = {"v": DECODER_VERSION, "slips": [], "others": [], "phash": None, "error": None}
    try:
        pages = load_pages(path)
        if not pages:
            entry["error"] = "no pages"
            return entry
        entry["phash"] = str(imagehash.phash(pages[0]))
        entry.update(decode_pages(pages))
    except Exception as e:  # one broken file must not stop the run
        entry["error"] = type(e).__name__  # message omitted: it often embeds the full file path
    return entry


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- Slip2Go

class Slip2Go:
    def __init__(self, base_url: str, secret: str):
        # accept a pasted endpoint URL too: keep only scheme + host
        self.base = re.sub(r"/api(/.*)?$", "", base_url.strip().rstrip("/"))
        self.secret = secret if secret.lower().startswith("bearer ") else f"Bearer {secret}"

    def _request(self, method: str, path: str, body=None) -> dict:
        req = urllib.request.Request(
            self.base + path, method=method,
            data=json.dumps(body).encode("utf-8") if body is not None else None,
            headers={"Authorization": self.secret, "Content-Type": "application/json",
                     "Accept": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=20) as r:
                raw = r.read()
        except urllib.error.HTTPError as e:
            raw = e.read()
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            return {"code": "NETWORK", "message": type(e).__name__}
        try:
            resp = json.loads(raw.decode("utf-8"))
            resp["code"] = str(resp.get("code", ""))
            return resp
        except (ValueError, AttributeError):
            return {"code": "BADRESPONSE", "message": "non-JSON response"}

    def call(self, method: str, path: str, body=None, tries: int = 3) -> dict:
        for attempt in range(tries):
            resp = self._request(method, path, body)
            if resp["code"] not in API_RETRY_CODES and not resp["code"].startswith("5"):
                return resp
            if attempt < tries - 1:
                time.sleep(1.5 * 2 ** attempt)
        return resp

    def account_info(self) -> dict:
        return self.call("GET", "/api/account/info")

    def verify_qr(self, payload: str, condition: dict) -> dict:
        body = {"payload": {"qrCode": payload, **({"checkCondition": condition} if condition else {})}}
        return self.call("POST", "/api/verify-slip/qr-code/info", body)


def api_condition(scfg: dict, rcfg: dict) -> dict:
    cond = {}
    if scfg.get("check_receiver", True):
        rec = {"accountType": "01" + str(rcfg.get("bank_code", "")),
               "accountNumber": re.sub(r"\D", "", str(rcfg.get("account_number", "")))}
        if rcfg.get("account_name_th"):
            rec["accountNameTH"] = rcfg["account_name_th"]
        if rcfg.get("account_name_en"):
            rec["accountNameEN"] = rcfg["account_name_en"]
        cond["checkReceiver"] = [rec]
    if scfg.get("check_date") and scfg.get("paid_after"):
        start = dt.datetime.combine(dt.date.fromisoformat(str(scfg["paid_after"])), dt.time(), THAI_TZ)
        cond["checkDate"] = {"type": "gte",
                             "date": start.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")}
    return cond


def is_retryable(code: str) -> bool:
    return code in API_RETRY_CODES or code.startswith("5")


def api_label(code: str | None) -> str:
    if code is None:
        return "pending"
    if code in API_OK_CODES:
        return "OK"
    if is_retryable(code):
        return "pending"
    return API_LABELS.get(code, f"ERROR {code}")


def slip_label(res: dict | None, rcfg: dict, check_rcv: bool) -> str:
    """api_label, but a Slip2Go receiver mismatch is re-judged on the bank data it returned.

    Slip2Go compares names strictly, and some banks (e.g. Krungsri) return only a cut-short
    English name even when the account number is exactly ours.
    """
    if not res:
        return "pending"
    label = api_label(res.get("code"))
    if label == "WRONG RECEIVER" and check_rcv and res.get("data") and check_receiver(res["data"], rcfg)[0] is True:
        return "OK"
    return label


_NAME_PREFIXES = ["เด็กชาย", "เด็กหญิง", "นางสาว", "น.ส.", "ด.ช.", "ด.ญ.", "นาง", "นาย", "บจก.", "บมจ.", "บริษัท"]
_NAME_WORDS = {"mr", "mrs", "ms", "miss", "mstr", "co", "ltd", "company", "limited", "plc", "จำกัด", "มหาชน"}


def normalize_name(s: str) -> str:
    s = nfc(s or "").lower().strip()
    s = re.sub(r"x{2,}|\*+", " ", s)  # masked parts
    changed = True
    while changed:
        changed = False
        for p in _NAME_PREFIXES:
            if s.startswith(p):
                s, changed = s[len(p):].lstrip(), True
    words = [w for w in re.split(r"[\s.,()\-]+", s) if w and w not in _NAME_WORDS]
    return re.sub(r"[^0-9a-z฀-๿]", "", "".join(words))


def name_matches(received: str, configured: str) -> str | None:
    """'exact', 'truncated' or None."""
    r, c = normalize_name(received), normalize_name(configured)
    if not r or not c:
        return None
    if r == c:
        return "exact"
    if (len(r) >= 4 and c.startswith(r)) or (len(c) >= 4 and r.startswith(c)):
        return "truncated"
    return None


def account_matches(received: str, configured: str) -> bool | None:
    """Compare a possibly masked account number (xxx-x-x5366-x) with the full one. None = can't tell."""
    rec = re.sub(r"[^0-9xX*]", "", received or "").lower().replace("*", "x")
    conf = re.sub(r"\D", "", configured or "")
    if not rec or not conf or sum(ch.isdigit() for ch in rec) < 4:
        return None
    if len(rec) == len(conf):
        return all(a == "x" or a == b for a, b in zip(rec, conf))
    pos = 0  # different length: visible digit groups must appear in order
    for group in re.findall(r"\d+", rec):
        i = conf.find(group, pos)
        if i < 0:
            return False
        pos = i + len(group)
    return True


def check_receiver(data: dict, rcfg: dict) -> tuple[bool | None, str, str, str]:
    """Local second check of where the money went. Returns (ok, text, receiver name, receiver account)."""
    recv = (data or {}).get("receiver") or {}
    acct = recv.get("account") or {}
    name = acct.get("name") or ""
    number = ((acct.get("bank") or {}).get("account")) or ((acct.get("proxy") or {}).get("account")) or ""
    bank_id = str((recv.get("bank") or {}).get("id") or "")
    bank_name = BANKS.get(bank_id, bank_id or "?")
    want_bank = str(rcfg.get("bank_code", ""))
    if want_bank and bank_id and bank_id != want_bank:
        return False, f"✗ wrong bank ({bank_name} {number})", name, number
    acc_ok = account_matches(number, str(rcfg.get("account_number", "")))
    if acc_ok is False:
        return False, f"✗ wrong account ({number})", name, number
    nm = [name_matches(name, rcfg.get(k, "")) for k in ("account_name_th", "account_name_en")]
    if not any(nm):
        return False, f"✗ name mismatch ({name})", name, number
    if acc_ok is None:  # bank hid the account number: name is all we can go on
        return None, f"? account not shown · name ok · {bank_name} {name}", name, number
    tag = "" if "exact" in nm else " (name truncated)"
    return True, f"✔ {bank_name} {number} · {name}{tag}", name, number


def thai_time(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return dt.datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(THAI_TZ).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return iso


def run_online(payloads: list[str], api_cache: dict, scfg: dict, rcfg: dict) -> dict:
    """Send new slip QRs to Slip2Go. Mutates api_cache. Returns run info for the summary."""
    info = {"sent": 0, "tokens": None, "stopped": ""}
    secret = os.environ.get("SLIP2GO_SECRET", "").strip()
    if not secret and SECRET_FILE.exists():
        secret = SECRET_FILE.read_text(encoding="utf-8-sig").strip()  # Notepad may add a BOM
    problems = []
    if not secret:
        problems.append(f"secret key (put it in {SECRET_FILE.name})")
    if not scfg.get("api_base_url"):
        problems.append("[slip2go] api_base_url")
    if scfg.get("check_receiver", True):
        if not re.sub(r"\D", "", str(rcfg.get("account_number", ""))):
            problems.append("[slip2go.receiver] account_number")
        if not rcfg.get("account_name_th") and not rcfg.get("account_name_en"):
            problems.append("[slip2go.receiver] account_name_th or account_name_en")
        if not rcfg.get("bank_code"):
            problems.append("[slip2go.receiver] bank_code")
    if problems:
        sys.exit("Slip2Go is enabled but these are missing: " + ", ".join(problems)
                 + ".\nFill them in, or set [slip2go] enabled = false.")

    client = Slip2Go(scfg["api_base_url"], secret)
    todo = [p for p in payloads if p not in api_cache or is_retryable(api_cache[p].get("code", ""))]
    cached = len(payloads) - len(todo)
    acc = client.account_info()
    if acc["code"].startswith("401"):
        info["stopped"] = API_AUTH_MESSAGES.get(acc["code"], acc.get("message", acc["code"]))
        print(f"Slip2Go: cannot connect: {info['stopped']}. Continuing offline.")
        return info
    d = acc.get("data") or {}
    info["tokens"] = d.get("tokenRemaining")
    quota = d.get("estimatedQuotaSlip")
    tok = "unknown" if info["tokens"] is None else f"{info['tokens']} (≈{quota} slips)" if quota is not None else info["tokens"]
    print(f"Slip2Go: {len(todo)} new slips to send ({cached} cached). Tokens remaining: {tok}")
    if not todo:
        return info
    cap = int(scfg.get("max_calls_per_run", 0) or 0)
    if cap and len(todo) > cap:
        print(f"  max_calls_per_run = {cap}: sending only {cap} now, the rest stay API_PENDING.")
        todo = todo[:cap]
    if isinstance(quota, (int, float)) and not d.get("checkSlipByCredit") and quota < len(todo):
        print(f"  Only about {int(quota)} slips left in your package: sending {int(quota)}.")
        todo = todo[:max(int(quota), 0)]
    if scfg.get("confirm_before_sending", True):
        if not sys.stdin.isatty():
            print("  Not sending: confirm_before_sending is on but there is no terminal to ask.")
            return info
        try:
            answer = input(f"  Send {len(todo)} slip QR codes to Slip2Go? [y/N] ")
        except EOFError:
            answer = ""
        if answer.strip().lower() not in ("y", "yes"):
            print("  Not sent.")
            return info
    cond = api_condition(scfg, rcfg)
    for p in tqdm(todo, desc="Checking slips with Slip2Go", unit="slip"):
        resp = client.verify_qr(p, cond)
        info["sent"] += 1
        code = resp["code"]
        if code.startswith("401"):
            info["stopped"] = API_AUTH_MESSAGES.get(code, resp.get("message", code))
            print(f"\nSlip2Go stopped: {info['stopped']}. Results so far are kept.")
            break
        api_cache[p] = {"code": code, "message": resp.get("message", ""), "data": resp.get("data"),
                        "checked_at": dt.datetime.now().isoformat(timespec="seconds")}
        try:
            save_json(API_CACHE_FILE, api_cache)  # every call costs money: never lose a result
        except OSError:
            pass  # still in memory and in slip2go.tmp; saved again below and recovered next run
    save_json(API_CACHE_FILE, api_cache)
    if info["sent"] and not info["stopped"]:
        after = client.account_info()
        if after["code"] == "200001":
            info["tokens"] = (after.get("data") or {}).get("tokenRemaining", info["tokens"])
    return info


# ---------------------------------------------------------------- zip handling

def zip_date(path: Path) -> dt.date:
    m = ZIP_DATE_RE.search(path.name)
    if m:
        try:
            return dt.date.fromisoformat(m.group(1))
        except ValueError:
            pass
    return dt.date.fromtimestamp(path.stat().st_mtime)


def pick_latest_zip() -> Path:
    zips = [p for p in INBOX.glob("*.zip") if p.is_file()]
    if not zips:
        sys.exit(f"No zip found in {INBOX}. Put payment-slips-YYYY-MM-DD.zip there.")
    return max(zips, key=lambda p: (zip_date(p), p.stat().st_mtime))


def member_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if not info.flag_bits & 0x800:  # no UTF-8 flag: zipfile decoded as cp437
        try:
            name = name.encode("cp437").decode("utf-8")
        except UnicodeError:
            pass
    return nfc(name.replace("\\", "/"))


def safe_rel_path(name: str) -> Path | None:
    parts = []
    for part in name.split("/"):
        if part in ("", ".", ".."):
            continue
        part = re.sub(r'[<>:"|?*\x00-\x1f]', "_", part).rstrip(" .")
        if part:
            parts.append(part)
    return Path(*parts) if parts else None


def extract_zip(zpath: Path) -> tuple[Path, dict[str, Path]]:
    """Extract into .cache/extracted/<zip stem>/ (once). Returns (dir, {member name: path})."""
    dest = EXTRACTED / zpath.stem
    marker = dest / ".extract_ok"
    stamp = f"{zpath.stat().st_size}:{zpath.stat().st_mtime_ns}"
    done = marker.exists() and marker.read_text(encoding="utf-8") == stamp
    members: dict[str, Path] = {}
    with zipfile.ZipFile(zpath) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = member_name(info)
            rel = safe_rel_path(name)
            if rel is None:
                continue
            target = dest / rel
            members[name] = target
            if not done:
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as src, open(target, "wb") as out:
                    shutil.copyfileobj(src, out)
    if not done:
        marker.write_text(stamp, encoding="utf-8")
    return dest, members


def read_index(members: dict[str, Path]) -> list[dict]:
    candidates = [n for n in members if n.rsplit("/", 1)[-1].lower() == "index.csv"]
    if not candidates:
        sys.exit("index.csv not found in the zip.")
    path = members[min(candidates, key=len)]
    text = path.read_bytes().decode("utf-8-sig")
    rows, seen = [], set()
    for raw in csv.DictReader(io.StringIO(text)):
        row = {nfc(k).strip(): nfc(v or "").strip() for k, v in raw.items() if k}
        lg = row.get("lgNumber", "")
        if not lg or lg in seen:
            continue
        seen.add(lg)
        rows.append(row)
    return rows


def match_files(rows: list[dict], members: dict[str, Path]) -> tuple[dict[str, tuple[str, Path]], list[str]]:
    """Map lgNumber -> (member name, path). Exact fileName match first, lgNumber prefix as fallback."""
    slip_members = {n: p for n, p in members.items()
                    if p.suffix.lower() in IMAGE_EXTS | {".pdf"}}
    by_base: dict[str, list[str]] = {}
    for n in slip_members:
        by_base.setdefault(n.rsplit("/", 1)[-1], []).append(n)
    by_base_fold = {b.casefold(): v for b, v in by_base.items()}
    used, out = set(), {}
    for row in rows:
        fname = nfc(row.get("fileName", "")).replace("\\", "/").rsplit("/", 1)[-1]
        names = by_base.get(fname) or by_base_fold.get(fname.casefold()) if fname else None
        if names:
            out[row["lgNumber"]] = (names[0], slip_members[names[0]])
            used.add(names[0])
    for row in rows:
        lg = row["lgNumber"]
        if lg in out:
            continue
        pref = [n for n in slip_members if n not in used and
                re.match(re.escape(lg) + r"(?:[_\-\s.]|$)", n.rsplit("/", 1)[-1], re.I)]
        if pref:
            name = sorted(pref)[0]
            out[lg] = (name, slip_members[name])
            used.add(name)
    unmatched = sorted(n for n in slip_members if n not in used)
    return out, unmatched


# ---------------------------------------------------------------- persistence

def load_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def save_json(path: Path, data) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    # On Windows the target can be briefly locked (antivirus, indexer, an editor): retry.
    for attempt in range(10):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            time.sleep(0.2 * (attempt + 1))
    os.replace(tmp, path)


def load_api_cache() -> dict:
    """Slip2Go results, plus any left in slip2go.tmp by a save that was interrupted."""
    cache = load_json(API_CACHE_FILE, {})
    leftover = load_json(API_CACHE_FILE.with_suffix(".tmp"), {})
    missing = {k: v for k, v in leftover.items() if k not in cache}
    if missing:
        cache.update(missing)
        print(f"Recovered {len(missing)} Slip2Go result(s) from an interrupted save.")
    return cache


def read_previous_workbook() -> tuple[dict[str, dict], object]:
    """Manual columns by lgNumber, and the passbook total, from the previous output."""
    if not XLSX.exists():
        return {}, None
    try:
        wb = openpyxl.load_workbook(XLSX, data_only=False)
    except Exception as e:
        sys.exit(f"Cannot read previous {XLSX.name} ({e}). Restore it from output/history/ and run again.")
    manual: dict[str, dict] = {}
    if "All" in wb.sheetnames:
        rows = wb["All"].iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows, [])]
        if "lgNumber" in header:
            idx = {c: header.index(c) for c in ["lgNumber", *MANUAL_COLS] if c in header}
            for r in rows:
                lg = r[idx["lgNumber"]] if idx["lgNumber"] < len(r) else None
                if not lg:
                    continue
                manual[str(lg).strip()] = {
                    c: ("" if (i := idx.get(c)) is None or i >= len(r) or r[i] is None else str(r[i]))
                    for c in MANUAL_COLS}
    passbook = None
    if "Summary" in wb.sheetnames:
        for a, b, *_ in wb["Summary"].iter_rows(values_only=True, max_col=2):
            if str(a or "").strip().lower() == PASSBOOK_LABEL:
                passbook = b
    return manual, passbook


def ensure_writable(path: Path) -> None:
    lock = path.with_name("~$" + path.name)
    busy = lock.exists()
    if path.exists() and not busy:
        try:
            with open(path, "r+b"):
                pass
        except PermissionError:
            busy = True
    if busy:
        sys.exit(f"{path.name} seems to be open in Excel. Close it and run again (nothing was changed).")


# ---------------------------------------------------------------- Excel output

def _set(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # data from participants is never a formula
    return cell


def row_fill(rec) -> PatternFill | None:
    kind = status_kind(rec["manual_status"])
    if rec["not_paid"]:
        return FILL["grey"]
    if kind == "verified":
        return FILL["green"]
    if rec["flag_codes"] or kind == "problem":
        return FILL["red"]
    if rec["in_sample"] == "yes":
        return FILL["yellow"]
    if rec["api_ok"]:
        return FILL["lightgreen"]
    return None


WIDTHS = {"lgNumber": 11, "firstName": 16, "lastName": 18, "fileStatus": 11, "file": 30, "slip_count": 6,
          "amount_paid": 10, "amount_detail": 14, "payment_status": 16, "receiver_check": 36,
          "transfer_time": 17, "sender_name": 20, "bank_name": 11, "transaction_ref": 26, "crc_ok": 7,
          "api_result": 14, "decode_method": 20, "flags": 40, "in_sample": 9, "first_seen": 11,
          "manual_status": 14, "checked_by": 12, "note": 40, "page": 5, "amount": 10,
          "receiver_name": 24, "receiver_account": 16, "api_code": 9, "api_message": 22, "checked_at": 19}


def write_table(ws, records, columns, manual_formulas=False, links=False):
    bold = Font(bold=True)
    for c, name in enumerate(columns, 1):
        ws.cell(row=1, column=c, value=name).font = bold
    all_col = {name: openpyxl.utils.get_column_letter(i) for i, name in enumerate(COLUMNS, 1)}
    for r, rec in enumerate(records, 2):
        fill = row_fill(rec) if "manual_status" in rec else None
        for c, name in enumerate(columns, 1):
            if manual_formulas and name in MANUAL_COLS:
                ref = f"INDEX(All!${all_col[name]}:${all_col[name]},MATCH($A{r},All!$A:$A,0))"
                cell = ws.cell(row=r, column=c, value=f'=IFERROR(IF({ref}="","",{ref}),"")')
            else:
                cell = _set(ws, r, c, rec.get(name, ""))
            if fill:
                cell.fill = fill
        if links and rec.get("_path"):
            try:
                rel = os.path.relpath(rec["_path"], OUTPUT).replace(os.sep, "/")
                ws.cell(row=r, column=columns.index("file") + 1).hyperlink = rel
            except ValueError:
                pass
    ws.freeze_panes = "A2"
    if records:
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(columns))}{len(records) + 1}"
    for c, name in enumerate(columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = WIDTHS.get(name, 14)


def build_workbook(records, to_check, slip_rows, changes, summary_rows, passbook, compare_label):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All"
    write_table(ws, records, COLUMNS, links=True)
    ms_col = openpyxl.utils.get_column_letter(COLUMNS.index("manual_status") + 1)
    dv = DataValidation(type="list", formula1=f'"{VERIFIED},{PROBLEM},{PENDING}"', allow_blank=True)
    dv.error, dv.errorTitle = "Pick a value from the list", "manual_status"
    ws.add_data_validation(dv)
    dv.add(f"{ms_col}2:{ms_col}{max(len(records) + 1, 2)}")

    ws = wb.create_sheet("To check")
    write_table(ws, to_check, COLUMNS, manual_formulas=True, links=True)
    ws.cell(row=len(to_check) + 3, column=1,
            value="manual_status / checked_by / note here mirror the All sheet. Edit them in the All sheet.").font = Font(italic=True)

    ws = wb.create_sheet("Slips")
    write_table(ws, slip_rows, SLIP_COLUMNS)

    ws = wb.create_sheet("Changes")
    cols = ["change", "lgNumber", "firstName", "lastName", "detail"]
    for c, name in enumerate(cols, 1):
        ws.cell(row=1, column=c, value=name).font = Font(bold=True)
    for r, ch in enumerate(changes, 2):
        for c, name in enumerate(cols, 1):
            _set(ws, r, c, ch.get(name, ""))
    ws.freeze_panes = "A2"
    for c, w in zip("ABCDE", (20, 11, 16, 18, 60)):
        ws.column_dimensions[c].width = w

    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34
    rows_by_label = {}
    r = 1
    for label, value in summary_rows:
        if label is None:
            r += 1
            continue
        a = ws.cell(row=r, column=1, value=label)
        rows_by_label[label] = r
        if value is None and label.isupper():
            a.font = Font(bold=True)
        if label == PASSBOOK_LABEL:
            _set(ws, r, 2, passbook)
            ws.cell(row=r, column=2).fill = FILL["yellow"]
            pb_row, cmp_row = r, rows_by_label[compare_label]
            r += 1
            ws.cell(row=r, column=1, value=f"difference (passbook − {compare_label.split(' (')[0]})")
            ws.cell(row=r, column=2,
                    value=f'=IF(OR(B{pb_row}="",NOT(ISNUMBER(B{cmp_row}))),"",B{pb_row}-B{cmp_row})')
        else:
            _set(ws, r, 2, value)
        r += 1
    return wb


# ---------------------------------------------------------------- main logic

def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description="Pre-screen LG26 payment slips.")
    ap.add_argument("--zip", type=Path, help="zip to process (default: latest in inbox/)")
    ap.add_argument("--redecode", action="store_true", help="ignore the decode cache")
    args = ap.parse_args()

    for d in (INBOX, OUTPUT, HISTORY, CACHE, EXTRACTED):
        d.mkdir(parents=True, exist_ok=True)
    if not (BASE / "config.toml").exists():
        sys.exit("config.toml not found. Copy config.example.toml to config.toml and fill it in.")
    cfg = tomllib.loads((BASE / "config.toml").read_text(encoding="utf-8"))
    fee = cfg.get("fee") or 0
    sample_size = int(cfg.get("sample_size", 40))
    sample_fraction = float(cfg.get("sample_fraction", 0.12))
    seed = cfg.get("seed", 2026)
    phash_max = int(cfg.get("phash_max_distance", 4))
    scfg = cfg.get("slip2go", {}) or {}
    rcfg = scfg.get("receiver", {}) or {}
    online = bool(scfg.get("enabled"))

    zpath = args.zip.resolve() if args.zip else pick_latest_zip()
    if not zpath.is_file():
        sys.exit(f"Zip not found: {zpath}")
    zdate = zip_date(zpath)
    run_time = dt.datetime.now()
    print(f"Using zip: {zpath.name}  (date {zdate.isoformat()})")

    ensure_writable(XLSX)
    ensure_writable(CSV_OUT)
    prev_manual, passbook = read_previous_workbook()
    state = load_json(STATE_FILE, {})
    prev_parts: dict[str, dict] = state.get("participants", {})
    first_run = not prev_parts

    _, members = extract_zip(zpath)
    rows = read_index(members)
    matched, unmatched_files = match_files(rows, members)

    # ---- decode (cached by file content hash)
    cache = {} if args.redecode else load_json(DECODE_CACHE_FILE, {})
    shas, decoded, errors = {}, {}, 0
    for lg, (_, path) in tqdm(sorted(matched.items()), desc="Decoding slips", unit="file"):
        try:
            sha = sha256_file(path)
        except OSError:
            continue
        shas[lg] = sha
        entry = cache.get(sha)
        if not entry or entry.get("v") != DECODER_VERSION:
            entry = cache[sha] = analyse_file(path)
        decoded[lg] = entry
        errors += bool(entry.get("error"))
    save_json(DECODE_CACHE_FILE, cache)

    # ---- online check (each unique valid slip QR at most once, ever)
    api_cache = load_api_cache()
    valid_payloads = sorted({s["payload"] for e in decoded.values() for s in e.get("slips", [])
                             if parse_slip(s["payload"])["crc_ok"]})
    api_info = {"sent": 0, "tokens": None, "stopped": ""}
    if online:
        api_info = run_online(valid_payloads, api_cache, scfg, rcfg)
    check_rcv = scfg.get("check_receiver", True) and bool(rcfg)

    # ---- per-participant records
    recs: dict[str, dict] = {}
    slip_rows: list[dict] = []
    for row in rows:
        lg = row["lgNumber"]
        prev = prev_parts.get(lg, {})
        uploaded = row.get("fileStatus", "").upper() == "UPLOADED"
        m = matched.get(lg)
        e = decoded.get(lg)
        rec = {
            "lgNumber": lg, "firstName": row.get("firstName", ""), "lastName": row.get("lastName", ""),
            "fileStatus": row.get("fileStatus", ""), "file": m[0].rsplit("/", 1)[-1] if m else "",
            "first_seen": prev.get("first_seen", ""), "_path": m[1] if m else None, "sha": shas.get(lg),
            "phash": e.get("phash") if e else None, "slips": [], "flag_list": [], "not_paid": False,
            "api_ok": False,
        }
        for s in (e or {}).get("slips", []):
            rec["slips"].append({**s, **parse_slip(s["payload"])})
        slips = rec["slips"]
        rec["slip_count"] = len(slips) if m else ""
        rec["bank_name"] = "; ".join(s["bank_name"] for s in slips)
        rec["transaction_ref"] = "; ".join(s["ref"] for s in slips)
        rec["crc_ok"] = "; ".join("yes" if s["crc_ok"] else "no" for s in slips)
        rec["decode_method"] = "; ".join(s["method"] for s in slips)

        if m and e:
            if e.get("error"):
                rec["decode_method"] = "unreadable: " + e["error"]
                rec["flag_list"].append("NO_QR")
            elif not slips and e.get("others"):
                rec["decode_method"] = "found: " + ", ".join(sorted({describe_other_qr(t) for t in e["others"]}))
                rec["flag_list"].append("NOT_SLIP_QR")
            elif not slips:
                rec["decode_method"] = "no QR after all retries"
                rec["flag_list"].append("NO_QR")
            if any(not s["crc_ok"] for s in slips):
                rec["flag_list"].append("BAD_CRC")
            if any(s["count"] > 1 for s in slips):
                rec["flag_list"].append("DUPLICATE_REF(same file)")
            if rec["sha"] and not rec["first_seen"]:
                rec["first_seen"] = zdate.isoformat()
        elif uploaded:
            rec["flag_list"].append("FILE_MISSING")
        else:
            rec["not_paid"] = True

        # API results per slip
        paid, details, api_texts, rcv_texts, times, senders = 0.0, [], [], [], [], []
        n_pending = n_counted = 0
        api_flags = set()
        for s in slips:
            res = api_cache.get(s["payload"]) if s["crc_ok"] else None
            code = res.get("code") if res else None
            label = slip_label(res, rcfg, check_rcv) if s["crc_ok"] else "not sent (bad CRC)"
            data = (res or {}).get("data") or {}
            amount = data.get("amount")
            rcv_ok, rcv_text, rcv_name, rcv_acct = (None, "", "", "")
            if data:
                rcv_ok, rcv_text, rcv_name, rcv_acct = check_receiver(data, rcfg) if check_rcv else (True, "not checked", "", "")
            if label == "pending":
                n_pending += 1
            counted = label == "OK" and rcv_ok is not False and isinstance(amount, (int, float))
            if label == "WRONG RECEIVER" or (label == "OK" and rcv_ok is False):
                api_flags.add("WRONG_RECEIVER")
            elif label == "FRAUD":
                api_flags.add("SLIP_FRAUD")
            elif label in ("NOT FOUND", "INVALID QR"):
                api_flags.add("SLIP_NOT_FOUND")
            elif label == "WRONG DATE":
                api_flags.add("WRONG_DATE")
            if counted:
                paid += amount
                n_counted += 1
            if isinstance(amount, (int, float)):
                details.append(baht(amount) + ("" if counted else "✗"))
            elif s["crc_ok"]:
                details.append("?")
            api_texts.append(label)
            if rcv_text:
                rcv_texts.append(rcv_text)
            if data.get("dateTime"):
                times.append(thai_time(data["dateTime"]))
            sender = ((data.get("sender") or {}).get("account") or {}).get("name")
            if sender:
                senders.append(sender)
            slip_rows.append({
                "lgNumber": lg, "firstName": rec["firstName"], "page": s["page"], "bank_name": s["bank_name"],
                "transaction_ref": s["ref"], "crc_ok": "yes" if s["crc_ok"] else "no",
                "amount": amount if isinstance(amount, (int, float)) else "",
                "transfer_time": thai_time(data.get("dateTime")), "sender_name": sender or "",
                "receiver_name": rcv_name, "receiver_account": rcv_acct, "receiver_check": rcv_text,
                "api_code": code or "", "api_message": (res or {}).get("message", "") if res else label,
                "checked_at": (res or {}).get("checked_at", ""),
            })
        has_api = any(t not in ("pending", "not sent (bad CRC)") for t in api_texts)
        rec["api_result"] = "; ".join(api_texts) if (has_api or online) and slips else ""
        rec["receiver_check"] = "; ".join(dict.fromkeys(rcv_texts))
        rec["transfer_time"] = "; ".join(times)
        rec["sender_name"] = "; ".join(dict.fromkeys(senders))
        rec["amount_detail"] = " + ".join(details) if has_api else ""
        rec["amount_paid"] = (int(paid) if paid.is_integer() else paid) if has_api else ""
        rec["flag_list"] += sorted(api_flags)
        if online and n_pending:
            rec["flag_list"].append("API_PENDING")
        # payment status: per participant, all slips together
        status = ""
        if slips and has_api:
            if n_pending:
                status = "⏳ partly checked"
            elif not n_counted:
                status = "✗ no valid slip"
            elif not fee:
                status = "set fee in config.toml"
            elif paid == fee:
                status = "✔ PAID"
            elif paid < fee:
                status = f"UNDERPAID −{baht(fee - paid)}"
                if scfg.get("check_amount", True):
                    rec["flag_list"].append(f"UNDERPAID(paid {baht(paid)} / fee {baht(fee)}, {len(slips)} slip{'s' * (len(slips) > 1)})")
            else:
                status = f"OVERPAID +{baht(paid - fee)}"
                if scfg.get("check_amount", True):
                    rec["flag_list"].append(f"OVERPAID(paid {baht(paid)} / fee {baht(fee)}, {len(slips)} slip{'s' * (len(slips) > 1)})")
        elif slips and online:
            status = "⏳ not checked"
        rec["payment_status"] = status
        rec["_counted_payloads"] = n_counted
        recs[lg] = rec

    def who(lg):
        r = recs[lg]
        return f"{lg} {r['firstName']}".strip()

    # duplicate transaction refs across participants
    by_ref: dict[str, list[str]] = {}
    for lg, r in recs.items():
        for ref in dict.fromkeys(s["ref"] for s in r["slips"]):
            by_ref.setdefault(ref, []).append(lg)
    for lgs in by_ref.values():
        if len(lgs) > 1:
            for lg in lgs:
                recs[lg]["flag_list"].append("DUPLICATE_REF(" + ", ".join(who(o) for o in lgs if o != lg) + ")")

    # duplicate / near-identical images. Slips from one bank app share a template, so a close
    # phash alone means nothing when both files carry different valid slip QRs (real data:
    # 623 such pairs, all different transfers). Near-identical only counts where a QR can't tell.
    hashed = [(lg, imagehash.hex_to_hash(r["phash"]), r["sha"], any(s["crc_ok"] for s in r["slips"]))
              for lg, r in recs.items() if r["phash"]]
    dup_img: dict[str, set[str]] = {}
    for i in range(len(hashed)):
        for j in range(i + 1, len(hashed)):
            a, ha, sa, va = hashed[i]
            b, hb, sb, vb = hashed[j]
            if sa == sb or (ha - hb <= phash_max and not (va and vb)):
                dup_img.setdefault(a, set()).add(b)
                dup_img.setdefault(b, set()).add(a)
    for lg, others in dup_img.items():
        recs[lg]["flag_list"].append("DUPLICATE_IMAGE(" + ", ".join(who(o) for o in sorted(others)) + ")")

    # ---- manual columns carried over; SLIP_CHANGED handling
    for lg, rec in recs.items():
        prev = prev_parts.get(lg, {})
        manual = prev_manual.get(lg) or prev.get("manual") or {}
        rec["manual_status"] = manual.get("manual_status", "")
        rec["checked_by"] = manual.get("checked_by", "")
        rec["note"] = manual.get("note", "")
        changed_pending = prev.get("changed_pending", False)
        if prev.get("sha") and rec["sha"] and prev["sha"] != rec["sha"]:
            changed_pending = True
            rec["changed_now"] = True
            rec["manual_status"] = PENDING
            msg = f"slip re-uploaded {zdate.isoformat()}"
            rec["note"] = f"{rec['note']} | {msg}" if rec["note"] else msg
        elif changed_pending and status_kind(rec["manual_status"]) != "pending":
            changed_pending = False  # treasurer has re-checked it
        rec["changed_pending"] = changed_pending
        if changed_pending:
            rec["flag_list"].append("SLIP_CHANGED")
        if not rec["manual_status"] and rec["sha"]:
            rec["manual_status"] = PENDING

    for rec in recs.values():
        rec["flag_list"].sort(key=lambda f: FLAG_ORDER.index(f.split("(")[0]))
        rec["flag_codes"] = [f.split("(")[0] for f in rec["flag_list"]]
        rec["flags"] = "; ".join(rec["flag_list"]) if rec["flag_list"] else (NOT_PAID if rec["not_paid"] else "")
        rec["api_ok"] = rec["payment_status"] == "✔ PAID" and not rec["flag_codes"]

    # ---- stable random sample of clean slips
    clean = sorted(lg for lg, r in recs.items()
                   if r["slips"] and all(s["crc_ok"] for s in r["slips"]) and not r["flag_codes"])
    kept = {lg for lg in recs if prev_parts.get(lg, {}).get("in_sample")}
    pool_prev = set(state.get("sample_pool", []))
    key = {lg: f"{lg}:{recs[lg]['sha']}" for lg in clean}
    new_pool = [lg for lg in clean if key[lg] not in pool_prev and lg not in kept]
    old_pool = [lg for lg in clean if key[lg] in pool_prev and lg not in kept]
    target = min(len(clean), max(sample_size, math.ceil(sample_fraction * len(clean))))
    need = max(target - len(kept & set(clean)), 0)
    want_new = 0 if first_run else round(sample_fraction * len(new_pool))
    rng = random.Random(f"{seed}|{zpath.name}")
    picks = rng.sample(new_pool, min(len(new_pool), max(need, want_new)))
    if need > len(picks):
        picks += rng.sample(old_pool, min(need - len(picks), len(old_pool)))
    sample = kept | set(picks)
    for lg, rec in recs.items():
        rec["in_sample"] = "yes" if lg in sample else ""
    new_pool_state = sorted(pool_prev | set(key.values()))

    # ---- sheets
    records = [recs[r["lgNumber"]] for r in rows]

    def check_group(r):
        if set(r["flag_codes"]) & set(MONEY_FLAGS):
            return 0
        if set(r["flag_codes"]) - {"API_PENDING"}:
            return 1
        if r["flag_codes"]:
            return 2
        return 3

    to_check = [r for r in records if r["flag_codes"]
                or (r["in_sample"] == "yes" and status_kind(r["manual_status"]) != "verified")]
    to_check.sort(key=lambda r: (check_group(r), status_kind(r["manual_status"]) == "verified"))

    changes = []
    if first_run:
        changes.append({"change": "first run", "detail": "Baseline created; changes are reported from the next run."})
    else:
        for lg, rec in recs.items():
            prev = prev_parts.get(lg)
            base = {"lgNumber": lg, "firstName": rec["firstName"], "lastName": rec["lastName"]}
            if prev is None:
                changes.append({**base, "change": "new in index.csv", "detail": rec["fileStatus"]})
            if rec.get("changed_now"):
                changes.append({**base, "change": "re-uploaded", "detail": "manual_status reset to Pending"})
            elif rec["sha"] and not (prev or {}).get("sha"):
                changes.append({**base, "change": "newly uploaded", "detail": rec["file"]})
            elif (prev or {}).get("sha") and not rec["sha"]:
                changes.append({**base, "change": "slip removed", "detail": rec["flags"]})
            old_codes = set((prev or {}).get("flags", []))
            added = [f for f in rec["flag_codes"] if f not in old_codes]
            cleared = [f for f in old_codes if f not in rec["flag_codes"]]
            if added:
                changes.append({**base, "change": "newly flagged", "detail": ", ".join(added)})
            if cleared:
                changes.append({**base, "change": "flag cleared", "detail": ", ".join(sorted(cleared))})
            if prev and rec["api_result"] and prev.get("api_result", "") != rec["api_result"]:
                changes.append({**base, "change": "API result changed",
                                "detail": f"{prev.get('api_result') or '(none)'} → {rec['api_result']}"})
            if prev and rec["payment_status"] == "✔ PAID" and "UNDERPAID" in prev.get("flags", []):
                changes.append({**base, "change": "payment completed", "detail": rec["amount_detail"]})
        for lg, prev in prev_parts.items():
            if lg not in recs and prev.get("in_index", True):
                m = prev.get("manual", {})
                changes.append({"lgNumber": lg, "firstName": prev.get("firstName", ""),
                                "lastName": prev.get("lastName", ""), "change": "removed from index.csv",
                                "detail": f"last manual_status: {m.get('manual_status', '')}"})
        if not changes:
            changes.append({"change": "no changes", "detail": f"Nothing changed since the previous run ({state.get('last_zip', '')})."})

    # ---- summary
    n_total = len(records)
    n_uploaded = sum(r["fileStatus"].upper() == "UPLOADED" for r in records)
    n_files = sum(bool(r["sha"]) for r in records)
    n_valid_people = sum(any(s["crc_ok"] for s in r["slips"]) for r in records)
    n_slips = sum(len(r["slips"]) for r in records)
    n_multi = sum(len(r["slips"]) > 1 for r in records)
    unique_refs = len({s["ref"] for r in records for s in r["slips"] if s["crc_ok"]})
    flag_counts = {f: sum(f in r["flag_codes"] for r in records) for f in FLAG_ORDER}
    n_not_paid = sum(r["not_paid"] for r in records)
    banks: dict[str, int] = {}
    for r in records:
        for s in r["slips"]:
            if s["crc_ok"]:
                banks[s["bank_name"]] = banks.get(s["bank_name"], 0) + 1
    expected = n_valid_people * fee if fee else "set fee in config.toml"
    ms = openpyxl.utils.get_column_letter(COLUMNS.index("manual_status") + 1)
    # money actually confirmed by Slip2Go, counting each slip once even if shared
    verified_total, seen_payloads = 0.0, set()
    for r in records:
        for s in r["slips"]:
            res = api_cache.get(s["payload"]) if s["crc_ok"] else None
            data = (res or {}).get("data") or {}
            if (s["payload"] in seen_payloads or not res or slip_label(res, rcfg, check_rcv) != "OK"
                    or not isinstance(data.get("amount"), (int, float))):
                continue
            if check_rcv and check_receiver(data, rcfg)[0] is False:
                continue
            seen_payloads.add(s["payload"])
            verified_total += data["amount"]
    any_api = bool(seen_payloads) or online
    labels = [slip_label(api_cache.get(p), rcfg, check_rcv) for p in valid_payloads]
    n_full = sum(r["payment_status"] == "✔ PAID" for r in records)
    compare_label = "verified amount total (THB)" if any_api else "expected total (THB)"
    summary_rows = [
        ("RUN", None),
        ("zip used", zpath.name), ("zip date", zdate.isoformat()), ("run time", run_time.strftime("%Y-%m-%d %H:%M:%S")),
        ("Slip2Go", "enabled" if online else "off (offline run)"),
        (None, None), ("PARTICIPANTS", None),
        ("total participants", n_total), ("uploaded (fileStatus UPLOADED)", n_uploaded),
        ("not uploaded (NOT_PAID)", n_not_paid), ("slip files found", n_files),
        ("files in zip not matched to index.csv", len(unmatched_files)),
        (None, None), ("SLIPS", None),
        ("participants with a valid slip QR", n_valid_people),
        ("slip QRs found (all files)", n_slips), ("participants with 2+ slips", n_multi),
        *[(f, flag_counts[f]) for f in FLAG_ORDER],
        ("rows flagged (any flag)", sum(bool(r["flag_codes"]) for r in records)),
        *[(f"bank: {b}", c) for b, c in sorted(banks.items())],
        (None, None), ("MONEY", None),
        ("unique valid transaction refs", unique_refs), ("fee (THB)", fee or "not set"),
        ("expected total (THB)", expected),
    ]
    if any_api:
        summary_rows += [
            ("participants fully paid (API)", n_full),
            ("underpaid / overpaid", f"{flag_counts['UNDERPAID']} / {flag_counts['OVERPAID']}"),
            ("verified amount total (THB)", verified_total),
        ]
    summary_rows += [(PASSBOOK_LABEL, "passbook")]
    if any_api:
        summary_rows += [
            (None, None), ("SLIP2GO", None),
            ("slips checked (cached total)", sum(lab != "pending" for lab in labels)),
            ("sent this run", api_info["sent"]),
            *[(f"result: {lab}", labels.count(lab)) for lab in sorted(set(labels))],
            ("tokens remaining", api_info["tokens"] if api_info["tokens"] is not None else "unknown (offline run)"),
        ]
        if api_info["stopped"]:
            summary_rows.append(("stopped because", api_info["stopped"]))
    summary_rows += [
        (None, None), ("CHECKING", None),
        ("spot-check sample size", len(sample)), ("rows in 'To check'", len(to_check)),
        ("manual: ✅ Verified", f'=COUNTIF(All!{ms}:{ms},"*Verified*")'),
        ("manual: ❌ Problem", f'=COUNTIF(All!{ms}:{ms},"*Problem*")'),
        ("manual: ⏳ Pending", f'=COUNTIF(All!{ms}:{ms},"*Pending*")'),
    ]
    if unmatched_files:
        summary_rows += [(None, None), ("UNMATCHED FILES IN ZIP", None)] + [("", n) for n in unmatched_files]

    # ---- write outputs (xlsx first; state only after outputs succeed)
    wb = build_workbook(records, to_check, slip_rows, changes, summary_rows, passbook, compare_label)
    if XLSX.exists():
        shutil.copy2(XLSX, HISTORY / f"slip_check_{run_time.strftime('%Y%m%d-%H%M%S')}.xlsx")
    tmp = XLSX.with_name("slip_check.tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, XLSX)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        sys.exit(f"{XLSX.name} is open in Excel. Close it and run again (state was not updated).")
    with open(CSV_OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in records:
            w.writerow([r[c] for c in COLUMNS])

    new_parts = {}
    for lg, prev in prev_parts.items():
        if lg not in recs:
            new_parts[lg] = {**prev, "in_index": False}
    for lg, r in recs.items():
        new_parts[lg] = {
            "sha": r["sha"], "first_seen": r["first_seen"], "in_sample": r["in_sample"] == "yes",
            "flags": r["flag_codes"], "changed_pending": r["changed_pending"], "in_index": True,
            "firstName": r["firstName"], "lastName": r["lastName"], "api_result": r["api_result"],
            "manual": {c: r[c] for c in MANUAL_COLS},
        }
    save_json(STATE_FILE, {"version": 1, "last_zip": zpath.name, "last_run": run_time.isoformat(timespec="seconds"),
                           "participants": new_parts, "sample_pool": new_pool_state})

    # ---- terminal summary (no names or refs here)
    print()
    line = f"Participants {n_total} · uploaded {n_uploaded} · not uploaded {n_not_paid} · files {n_files}"
    if any_api:
        line += f" · fully paid {n_full} · underpaid {flag_counts['UNDERPAID']} · overpaid {flag_counts['OVERPAID']}"
    print(line)
    print(f"Slip QRs: {n_slips} in {n_valid_people} participants ({n_multi} with 2+ slips), unique valid refs {unique_refs}")
    print("Flags: " + (", ".join(f"{f}={c}" for f, c in flag_counts.items() if c) or "none"))
    if errors:
        print(f"Unreadable files: {errors} (see decode_method column)")
    if unmatched_files:
        print(f"Files in zip not matched to index.csv: {len(unmatched_files)}")
    print(f"Expected total: {expected if isinstance(expected, str) else f'{baht(expected)} THB'}"
          + (f" · verified by Slip2Go: {baht(verified_total)} THB" if any_api else ""))
    print(f"Spot-check sample: {len(sample)}   To check: {len(to_check)} rows")
    real_changes = 0 if first_run or changes[0]["change"] == "no changes" else len(changes)
    print(f"Changes since previous run: {'first run' if first_run else real_changes}")
    print(f"Wrote {XLSX.relative_to(BASE)} and {CSV_OUT.relative_to(BASE)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
